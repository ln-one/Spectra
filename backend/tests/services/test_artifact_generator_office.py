from pathlib import Path

import pytest

from services.artifact_generator import office as office_module
from services.artifact_generator.office import ArtifactOfficeMixin


class _OfficeGenerator(ArtifactOfficeMixin):
    def __init__(self, base_dir: Path):
        self.base_dir = base_dir

    def get_storage_path(self, project_id: str, artifact_type: str, artifact_id: str):
        ext = artifact_type
        target_dir = self.base_dir / project_id / artifact_type
        target_dir.mkdir(parents=True, exist_ok=True)
        return str(target_dir / f"{artifact_id}.{ext}")


@pytest.mark.asyncio
async def test_generate_pptx_creates_openable_file(tmp_path: Path):
    generator = _OfficeGenerator(tmp_path)
    output = await generator.generate_pptx(
        {"title": "Demo PPTX", "slides": [{"title": "S1", "content": "C1"}]},
        "project-1",
        "artifact-1",
    )

    path = Path(output)
    assert path.exists()
    assert path.suffix == ".pptx"
    assert path.stat().st_size > 0

    from pptx import Presentation

    presentation = Presentation(str(path))
    assert len(presentation.slides) >= 1


@pytest.mark.asyncio
async def test_generate_docx_creates_openable_file(tmp_path: Path):
    generator = _OfficeGenerator(tmp_path)
    output = await generator.generate_docx(
        {
            "title": "Demo DOCX",
            "sections": [{"title": "Intro", "content": "Hello"}],
        },
        "project-1",
        "artifact-2",
    )

    path = Path(output)
    assert path.exists()
    assert path.suffix == ".docx"
    assert path.stat().st_size > 0

    from docx import Document

    document = Document(str(path))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert any("Demo DOCX" in text for text in texts)
    assert any("Intro" in text for text in texts)


@pytest.mark.asyncio
async def test_generate_pptx_fails_explicitly_when_placeholder_disabled(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    generator = _OfficeGenerator(tmp_path)
    monkeypatch.delenv("ALLOW_OFFICE_PLACEHOLDER_ARTIFACTS", raising=False)

    async def _render_fail(*args, **kwargs):
        return False

    monkeypatch.setattr(generator, "_render_pptx_with_marp", _render_fail)
    monkeypatch.setattr(office_module, "Presentation", None)

    with pytest.raises(RuntimeError, match="placeholder artifacts are disabled"):
        await generator.generate_pptx(
            {"title": "Demo PPTX", "slides": [{"title": "S1", "content": "C1"}]},
            "project-1",
            "artifact-disabled",
        )
